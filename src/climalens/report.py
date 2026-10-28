"""The ORSA climate brief as a Word document (.docx), built from the same text and numbers as the web app."""

import io
from datetime import date

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

GREEN = RGBColor(0x16, 0x78, 0x4F)
GREY = RGBColor(0x5F, 0x6B, 0x66)


def _pct(v: float) -> str:
    return f"{v:.1f} %"


def _bn(v: float) -> str:
    return f"EUR {v / 1e9:.2f} bn"


def brief_text(t: pd.DataFrame, a: dict, inp) -> str:
    """The brief as Markdown: t = results table with the top hotspot per scenario, a = assumptions, inp = munichre.Inputs."""
    cp, nz = t.iloc[0], t.iloc[-1]
    invested = inp.portfolio["carrying_amount_meur"].sum() / 1e3
    emissions = (inp.portfolio["scope1_financed_tco2e"] + inp.portfolio["scope2_financed_tco2e"]).sum() / 1e6
    return f"""### Scope
This section assesses Munich Re's exposure to climate change over the ORSA horizon 2026–2030 on both sides of the balance sheet. It covers corporate equity, loans and bonds with a carrying amount of EUR {invested:.1f} bn and financed Scope 1 and 2 emissions of {emissions:.1f} Mt CO2e, and the natural catastrophe business with average net large losses of {_bn(inp.baseline_eur)} per year (2021–2025).

### Results
The combined burden ranges from {_pct(cp.net_pct_equity)} of group equity under {cp.scenario} to {_pct(nz.net_pct_equity)} under {nz.scenario}. Under {nz.scenario} a carbon price of about {nz.carbon_price_eur:.0f} EUR/t reduces the value of the investments by {_bn(nz.asset_loss_eur)} ({_pct(nz.asset_pct_equity)} of equity). Under {cp.scenario}, extra natural catastrophe losses after repricing reach {_bn(cp.liability_extra_eur)} ({_pct(cp.liability_pct_equity)} of equity) over five years.

### Assessment
Up to 2030 the two risks do not offset each other. Transition risk on the asset side dominates in ambitious policy scenarios, while physical risk grows slowly and differs little between scenarios, because warming up to 2030 is largely locked in. The asset-side risk is concentrated in a few carbon-intensive industries, above all power generation. On the underwriting side, the largest single source of extra losses is {cp.top_hotspot} ({cp.top_hotspot_share:.0f} % of the extra nat-cat losses under {cp.scenario}).

### Key assumptions
{a['pass_through']:.0%} of carbon costs passed on to customers; valuation multiple {a['ev_ebitda_multiple']:g}×; loans and bonds take {a['debt_loss_factor']:.0%} of the shareholder loss; {a['repricing']:.0%} of extra nat-cat losses recovered through repricing; nat-cat trend {a['trend']:.1%} per year. Physical damage to investee companies, earthquakes (not climate-driven) and effects after 2030 are not included.
"""


def _shade(cell, hex_fill: str):
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), hex_fill)
    props.append(shading)


def brief_docx(markdown: str, table: pd.DataFrame, as_of: date | None = None) -> bytes:
    """Word file with a title, a results table and the brief text.

    markdown: the brief as produced by the app ("### Heading" lines and paragraphs).
    table: one row per scenario with scenario, carbon_price_eur and the *_pct_equity columns.
    """
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    doc.styles["Heading 2"].font.color.rgb = GREEN

    title = doc.add_heading("ORSA climate brief: Munich Re", level=0)
    title.runs[0].font.color.rgb = GREEN
    sub = doc.add_paragraph()
    run = sub.add_run(f"ClimaLens draft · ORSA horizon 2026–2030 · {(as_of or date.today()):%d %B %Y}")
    run.font.color.rgb = GREY
    run.font.size = Pt(10)

    doc.add_heading("Results at a glance", level=2)
    head = ["NGFS scenario", "Carbon price 2030 (EUR/t)", "Investments (% of equity)",
            "Nat-cat (% of equity)", "Total (% of equity)"]
    grid = doc.add_table(rows=1, cols=len(head))
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(grid.rows[0].cells, head):
        cell.text = text
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(cell, "16784F")
    for r in table.itertuples():
        values = [r.scenario, f"{r.carbon_price_eur:.0f}", f"{r.asset_pct_equity:.1f}",
                  f"{r.liability_pct_equity:.1f}", f"{r.net_pct_equity:.1f}"]
        cells = grid.add_row().cells
        for i, (cell, text) in enumerate(zip(cells, values)):
            cell.text = text
            if i:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        cells[-1].paragraphs[0].runs[0].font.bold = True

    for block in markdown.strip().split("\n\n"):
        block = block.strip()
        if block.startswith("### "):
            heading, _, rest = block.partition("\n")
            doc.add_heading(heading[4:].strip(), level=2)
            block = rest.strip()
        if block:
            doc.add_paragraph(block)

    note = doc.add_paragraph()
    run = note.add_run("Numbers computed by the ClimaLens stress test (github.com/BoyensKjell/ClimaLens). "
                       "Sources: Munich Re annual reports 2021–2025; NGFS Phase V; Swiss Re Institute sigma 4/2021.")
    run.font.size = Pt(9)
    run.font.color.rgb = GREY

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()

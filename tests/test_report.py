import io

from docx import Document

from climalens import munichre
from climalens.report import brief_docx


def test_brief_is_a_word_file_with_headings_and_results():
    inp = munichre.load_inputs()
    t = munichre.trap(inp)
    t["carbon_price_eur"] = [p[0].carbon_price_eur for p in inp.pairs]
    data = brief_docx("### Scope\nFirst paragraph.\n\n### Results\nSecond paragraph.", t)
    doc = Document(io.BytesIO(data))
    text = [p.text for p in doc.paragraphs]
    assert "Scope" in text and "Results" in text and "Second paragraph." in text
    rows = doc.tables[0].rows
    assert len(rows) == 1 + len(t)
    assert rows[-1].cells[-1].text == f"{t['net_pct_equity'].iloc[-1]:.1f}"
